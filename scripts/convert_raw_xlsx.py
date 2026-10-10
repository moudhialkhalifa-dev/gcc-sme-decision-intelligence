"""Convert the UCI Online Retail II workbook (two sheets) into one raw CSV, unchanged values."""
import csv, sys, hashlib
from openpyxl import load_workbook

src, dst = sys.argv[1], sys.argv[2]
print("sha256", hashlib.sha256(open(src, "rb").read()).hexdigest())
wb = load_workbook(src, read_only=True)
total = 0
with open(dst, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    for i, ws in enumerate(wb.worksheets):
        n = 0
        for j, row in enumerate(ws.iter_rows(values_only=True)):
            if j == 0:
                print(ws.title, row)
                if i == 0:
                    w.writerow(row)
                continue
            w.writerow(["" if v is None else v for v in row])
            n += 1
        print(ws.title, n)
        total += n
print("total", total)
